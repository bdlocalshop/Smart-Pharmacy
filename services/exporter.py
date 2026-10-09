import os
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database.db import get_connection
from models.reports_repo import get_expiry_alerts
from models.sales_repo import get_sales_history, get_sales_summary, get_top_selling_medicines

def get_exports_dir():
    """Returns absolute path to the exports directory."""
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    exports_dir = os.path.join(base_dir, "exports")
    os.makedirs(exports_dir, exist_ok=True)
    return exports_dir

def export_inventory_to_excel(target_path=None):
    """Exports full inventory with batches, cost, and shelf location to formatted Excel."""
    if not target_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_path = os.path.join(get_exports_dir(), f"pharmacy_inventory_{timestamp}.xlsx")

    conn = get_connection()
    cursor = conn.cursor()
    query = """
        SELECT 
            m.name as medicine_name,
            m.generic_name,
            m.company,
            m.category,
            m.shelf_location,
            b.batch_number,
            COALESCE(s.name, 'N/A') as supplier_name,
            b.purchase_date,
            b.expiry_date,
            b.purchase_price,
            b.selling_price,
            b.current_qty,
            (b.current_qty * b.purchase_price) as total_cost,
            (b.current_qty * b.selling_price) as total_retail,
            CASE 
                WHEN b.expiry_date < DATE('now') THEN 'Expired'
                WHEN b.expiry_date <= DATE('now', '+30 days') THEN 'Expiring Soon'
                ELSE 'Valid'
            END as status
        FROM batches b
        JOIN medicines m ON b.medicine_id = m.id
        LEFT JOIN suppliers s ON b.supplier_id = s.id
        ORDER BY m.name ASC, b.expiry_date ASC;
    """
    cursor.execute(query)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory Report"

    headers = [
        "Medicine Name", "Generic Name", "Company", "Category", "Shelf Location",
        "Batch #", "Supplier / Source", "Purchase Date", "Expiry Date",
        "Cost Price", "Selling Price", "Stock Qty", "Cost Value", "Retail Value", "Status"
    ]
    ws.append(headers)

    # Style Header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for col_num, col in enumerate(ws.iter_cols(min_row=1, max_row=1, min_col=1, max_col=len(headers)), 1):
        for cell in col:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data Rows
    for row in rows:
        ws.append(list(row))

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(target_path)
    return target_path

def export_sales_to_excel(start_date=None, end_date=None, target_path=None):
    """Exports sales transactions report to formatted Excel."""
    if not target_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_path = os.path.join(get_exports_dir(), f"pharmacy_sales_{timestamp}.xlsx")

    sales = get_sales_history(start_date=start_date, end_date=end_date, limit=10000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales History"

    headers = [
        "Invoice No", "Date & Time", "Customer Name", "Customer Phone",
        "Items Sold", "Subtotal", "Discount", "Grand Total", "Payment Method", "Estimated Profit"
    ]
    ws.append(headers)

    # Header styling
    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for s in sales:
        ws.append([
            s["invoice_no"], s["sale_date"], s["customer_name"], s["customer_phone"],
            s["item_count"], s["subtotal"], s["discount"], s["grand_total"],
            s["payment_method"], round(s["estimated_profit"], 2)
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(target_path)
    return target_path

def generate_sales_text_report(start_date=None, end_date=None, period_label="Selected Period", target_path=None):
    """
    Generates a beautifully formatted professional text report (.txt) of sales,
    KPIs, payment method breakdown, top-selling products, and itemized bills.
    """
    if not target_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        slug = period_label.lower().replace(" ", "_").replace("(", "").replace(")", "")
        target_path = os.path.join(get_exports_dir(), f"pharmacy_sales_report_{slug}_{timestamp}.txt")

    summary = get_sales_summary(start_date=start_date, end_date=end_date)
    sales = get_sales_history(start_date=start_date, end_date=end_date, limit=10000)
    top_meds = get_top_selling_medicines(start_date=start_date, end_date=end_date, limit=10)

    # Format Date Range string
    if start_date and end_date:
        if start_date == end_date:
            date_range_str = f"{start_date} ({period_label})"
        else:
            date_range_str = f"{start_date} to {end_date} ({period_label})"
    elif start_date:
        date_range_str = f"From {start_date} onwards ({period_label})"
    else:
        date_range_str = "All Time (All Recorded Sales)"

    total_inv = summary["total_invoices"]
    total_units = summary["total_items_sold"]
    gross_sub = summary.get("gross_subtotal", summary["total_revenue"] + summary["total_discount"])
    discount = summary["total_discount"]
    net_rev = summary["total_revenue"]
    cogs = summary.get("total_cogs", 0.0)
    est_profit = summary["estimated_gross_profit"]
    margin_pct = (est_profit / net_rev * 100) if net_rev > 0 else 0.0
    avg_ticket = (net_rev / total_inv) if total_inv > 0 else 0.0

    gen_time_str = datetime.now().strftime("%Y-%m-%d %I:%M:%S %p")
    w = 88
    div_eq = "=" * w
    div_dash = "-" * w

    lines = []
    lines.append(div_eq)
    lines.append("SMART PHARMACY MANAGEMENT SYSTEM".center(w))
    lines.append("SALES & PERFORMANCE AUDIT REPORT".center(w))
    lines.append(div_eq)
    lines.append("")
    lines.append(" REPORT METADATA")
    lines.append(f"   Generated On       : {gen_time_str}")
    lines.append(f"   Report Period      : {date_range_str}")
    lines.append("   Store / Facility   : Smart Pharmacy Retail Hub")
    lines.append("   System Engine      : SQLite3 Enterprise POS Engine")
    lines.append("")
    lines.append(div_dash)
    lines.append(" KEY PERFORMANCE INDICATORS (KPIs)")
    lines.append(div_dash)
    lines.append(f"   Total Invoices Issued       : {total_inv:>12,d} bills")
    lines.append(f"   Total Units Sold            : {total_units:>12,d} units")
    lines.append(f"   Gross Sales (Subtotal)      : ${gross_sub:>12,.2f}")
    lines.append(f"   Total Discounts Given       : ${discount:>12,.2f}")
    lines.append(f"   Net Sales Revenue           : ${net_rev:>12,.2f}")
    lines.append(f"   Cost of Goods Sold (COGS)   : ${cogs:>12,.2f}")
    lines.append(f"   Estimated Gross Profit      : ${est_profit:>12,.2f}")
    lines.append(f"   Gross Profit Margin         : {margin_pct:>12.2f} %")
    lines.append(f"   Average Invoice Value       : ${avg_ticket:>12,.2f}")
    lines.append("")
    lines.append(div_dash)
    lines.append(" PAYMENT METHOD BREAKDOWN")
    lines.append(div_dash)
    lines.append(f"   {'Payment Mode':<28} {'Transactions':>14} {'Total Amount ($)':>18} {'Share (%)':>14}")
    lines.append("   " + "-" * 78)

    pm_list = summary.get("payment_breakdown", [])
    if pm_list:
        for pm in pm_list:
            share = (pm["amount"] / net_rev * 100) if net_rev > 0 else 0.0
            lines.append(f"   {pm['payment_method']:<28} {pm['count']:>14,d} ${pm['amount']:>17,.2f} {share:>13.2f}%")
    else:
        lines.append("   (No transactions recorded for this period)")
    lines.append("   " + "-" * 78)
    lines.append(f"   {'TOTAL':<28} {total_inv:>14,d} ${net_rev:>17,.2f} {'100.00%':>14}")
    lines.append("")

    lines.append(div_dash)
    lines.append(" TOP SELLING MEDICINES IN THIS PERIOD")
    lines.append(div_dash)
    lines.append(f"   {'#':<3} {'Medicine Name':<24} {'Manufacturer':<22} {'Qty Sold':>9} {'Revenue ($)':>12} {'Profit ($)':>12}")
    lines.append("   " + "-" * 88)
    if top_meds:
        for idx, tm in enumerate(top_meds, 1):
            name_d = tm["medicine_name"][:23]
            comp_d = tm["company"][:21]
            lines.append(f"   {idx:<3} {name_d:<24} {comp_d:<22} {tm['units_sold']:>9,d} ${tm['total_sales']:>11,.2f} ${tm['profit']:>11,.2f}")
    else:
        lines.append("   (No product sales recorded in this period)")
    lines.append("")

    lines.append(div_dash)
    lines.append(" ITEMIZED TRANSACTION ARCHIVE")
    lines.append(div_dash)
    lines.append(f"   {'Invoice #':<16} {'Date & Time':<17} {'Patient / Customer':<20} {'Items':>6} {'Total ($)':>10} {'Profit ($)':>10}")
    lines.append("   " + "-" * 88)
    if sales:
        for s in sales:
            inv = s["invoice_no"]
            dt = s["sale_date"][:16] if s.get("sale_date") else "-"
            cust = (s.get("customer_name") or "Walk-in")[:19]
            lines.append(f"   {inv:<16} {dt:<17} {cust:<20} {s['item_count']:>6} ${s['grand_total']:>9,.2f} ${s['estimated_profit']:>9,.2f}")
    else:
        lines.append("   (No invoices recorded in this period)")

    lines.append("")
    lines.append(div_eq)
    lines.append("*** END OF REPORT ***".center(w))
    lines.append(div_eq)

    report_content = "\n".join(lines)
    with open(target_path, "w", encoding="utf-8") as f:
        f.write(report_content)

    return target_path

def export_sales_reports(start_date=None, end_date=None, period_label="Today"):
    """
    Exports both Excel spreadsheet (.xlsx) and formatted Text audit report (.txt)
    for the selected period.
    Returns (excel_path, txt_path).
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    slug = period_label.lower().replace(" ", "_").replace("(", "").replace(")", "")

    excel_filename = f"pharmacy_sales_{slug}_{timestamp}.xlsx"
    excel_path = os.path.join(get_exports_dir(), excel_filename)
    export_sales_to_excel(start_date=start_date, end_date=end_date, target_path=excel_path)

    txt_filename = f"pharmacy_sales_report_{slug}_{timestamp}.txt"
    txt_path = os.path.join(get_exports_dir(), txt_filename)
    generate_sales_text_report(start_date=start_date, end_date=end_date, period_label=period_label, target_path=txt_path)

    return excel_path, txt_path

def export_expiry_to_excel(target_path=None):
    """Exports list of expired and near-expiry stock for supplier returns."""
    if not target_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_path = os.path.join(get_exports_dir(), f"pharmacy_expiry_radar_{timestamp}.xlsx")

    alerts = get_expiry_alerts()

    wb = Workbook()
    ws = wb.active
    ws.title = "Expiry Radar"

    headers = [
        "Status", "Days Left", "Medicine Name", "Generic Name", "Company",
        "Shelf Location", "Batch #", "Supplier / Source", "Expiry Date",
        "Quantity", "Cost Price", "Cost Value Locked"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for a in alerts:
        ws.append([
            a["alert_level"], a["days_remaining"], a["medicine_name"], a["generic_name"],
            a["company"], a["shelf_location"], a["batch_number"], a["supplier_name"],
            a["expiry_date"], a["current_qty"], a["purchase_price"], round(a["cost_value"], 2)
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(target_path)
    return target_path

def generate_thermal_receipt_text(sale_details, shop_name="SMART PHARMACY", shop_phone="+880 1711-000000", shop_address="12 Central Road, Healthcare Plaza"):
    """
    Generates a clean 80mm thermal printer style text receipt.
    Ideal for direct printer spooling or on-screen preview.
    """
    lines = []
    width = 44
    lines.append("=" * width)
    lines.append(shop_name.center(width))
    lines.append(shop_address.center(width))
    lines.append(f"Phone: {shop_phone}".center(width))
    lines.append("-" * width)
    lines.append(f"Invoice: {sale_details['invoice_no']}".ljust(width))
    lines.append(f"Date   : {sale_details['sale_date']}".ljust(width))
    lines.append(f"Patient: {sale_details.get('customer_name') or 'Walk-in'}".ljust(width))
    if sale_details.get('customer_phone'):
        lines.append(f"Contact: {sale_details['customer_phone']}".ljust(width))
    lines.append("-" * width)
    lines.append(f"{'Item':<22}{'Qty':>4} {'Price':>8} {'Total':>8}")
    lines.append("-" * width)

    for itm in sale_details.get("items", []):
        med_display = itm['medicine_name'][:20]
        lines.append(f"{med_display:<22}{itm['quantity']:>4} {itm['unit_price']:>8.2f} {itm['total_price']:>8.2f}")
        # Secondary line for batch & expiry
        batch_line = f"  (Batch: {itm['batch_number']} | Exp: {itm['expiry_date']})"
        lines.append(batch_line[:width])

    lines.append("-" * width)
    lines.append(f"{'Subtotal:':<30}{sale_details['subtotal']:>14.2f}")
    if sale_details.get("discount", 0) > 0:
        lines.append(f"{'Discount:':<30}{sale_details['discount']:>14.2f}")
    lines.append(f"{'GRAND TOTAL:':<30}{sale_details['grand_total']:>14.2f}")
    lines.append(f"{'Payment Method:':<30}{sale_details.get('payment_method', 'Cash'):>14}")
    lines.append("=" * width)
    lines.append("Thank you for your visit!".center(width))
    lines.append("Prescription medicines are non-refundable.".center(width))
    lines.append("=" * width)
    return "\n".join(lines)
